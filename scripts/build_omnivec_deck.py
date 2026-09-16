"""Generate the OmniVec value-focused pitch deck.

Audience: field engineers pitching OmniVec to customer engineers.
Emphasis: what it offers, use cases, why choose it. Architecture kept light.

Run: python scripts/build_omnivec_deck.py
Output: OmniVec-Overview.pptx (repo root)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
NAVY   = RGBColor(0x0B, 0x25, 0x45)
NAVY2  = RGBColor(0x13, 0x33, 0x5C)
AZURE  = RGBColor(0x2B, 0x7F, 0xFF)
CYAN   = RGBColor(0x22, 0xC3, 0xE6)
TEAL   = RGBColor(0x1B, 0xC9, 0xA5)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x14, 0x1B, 0x2E)
SLATE  = RGBColor(0x53, 0x60, 0x76)
LIGHT  = RGBColor(0xF4, 0xF7, 0xFC)
LIGHT2 = RGBColor(0xE7, 0xEE, 0xF8)
AMBER  = RGBColor(0xF5, 0xA6, 0x23)
ROSE   = RGBColor(0xE0, 0x6A, 0x6A)
MUTE   = RGBColor(0x9A, 0xA6, 0xB8)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def rect(s, x, y, w, h, color, line=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.05):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, col, bold, ital) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(sz)
            r.font.color.rgb = col
            r.font.bold = bold
            r.font.italic = ital
            r.font.name = "Segoe UI"
    return tb


def bullets(s, x, y, w, h, items, size=16, color=INK, gap=9, marker="\u2014",
            marker_color=AZURE, line_spacing=1.06):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        mk = p.add_run()
        mk.text = marker + "  "
        mk.font.size = Pt(size)
        mk.font.bold = True
        mk.font.color.rgb = marker_color
        mk.font.name = "Segoe UI"
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True
            r1.font.color.rgb = color; r1.font.name = "Segoe UI"
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.color.rgb = SLATE
            r2.font.name = "Segoe UI"
        else:
            r = p.add_run(); r.text = item
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.name = "Segoe UI"
    return tb


def content_header(s, kicker, title, idx):
    bg(s, WHITE)
    rect(s, Inches(0), Inches(0), Inches(0.16), SH, AZURE)
    txt(s, Inches(0.7), Inches(0.42), Inches(11.5), Inches(0.4),
        [[(kicker.upper(), 12, AZURE, True, False)]], space_after=0)
    txt(s, Inches(0.68), Inches(0.72), Inches(12.0), Inches(0.9),
        [[(title, 30, INK, True, False)]], space_after=0)
    rect(s, Inches(0.72), Inches(1.42), Inches(1.1), Inches(0.06), CYAN)
    footer(s, idx)
    return Inches(1.75)


def footer(s, idx):
    txt(s, Inches(0.7), Inches(7.02), Inches(9), Inches(0.35),
        [[("OmniVec  \u00b7  Universal Vector Ingestion Platform", 9, SLATE, False, False)]],
        space_after=0)
    txt(s, Inches(11.4), Inches(7.02), Inches(1.4), Inches(0.35),
        [[(str(idx), 9, SLATE, False, False)]], align=PP_ALIGN.RIGHT, space_after=0)


def card(s, x, y, w, h, title, body, accent=AZURE, title_size=15, body_size=12.5,
         fill=LIGHT, tcol=INK, bcol=SLATE):
    panel = rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    panel.adjustments[0] = 0.06
    rect(s, x, y, Inches(0.09), h, accent)
    pad = Inches(0.22)
    tb = s.shapes.add_textbox(x + pad, y + Inches(0.14), w - pad - Inches(0.15), h - Inches(0.24))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]; p.space_after = Pt(5)
    r = p.add_run(); r.text = title; r.font.size = Pt(title_size)
    r.font.bold = True; r.font.color.rgb = tcol; r.font.name = "Segoe UI"
    for line in body:
        bp = tf.add_paragraph(); bp.space_after = Pt(3); bp.line_spacing = 1.05
        rr = bp.add_run(); rr.text = line; rr.font.size = Pt(body_size)
        rr.font.color.rgb = bcol; rr.font.name = "Segoe UI"
    return panel


def icon_card(s, x, y, w, h, glyph, title, body, accent):
    """Value card: colored glyph chip + title + one/two lines."""
    panel = rect(s, x, y, w, h, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    panel.adjustments[0] = 0.06
    chip = rect(s, x + Inches(0.22), y + Inches(0.22), Inches(0.52), Inches(0.52), accent,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    chip.adjustments[0] = 0.25
    tf = chip.text_frame; tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = glyph; r.font.size = Pt(20); r.font.bold = True
    r.font.color.rgb = WHITE; r.font.name = "Segoe UI"
    txt(s, x + Inches(0.9), y + Inches(0.24), w - Inches(1.05), Inches(0.5),
        [[(title, 15, INK, True, False)]], space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    tb = s.shapes.add_textbox(x + Inches(0.24), y + Inches(0.92), w - Inches(0.45), h - Inches(1.0))
    tf2 = tb.text_frame; tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = Emu(0); tf2.margin_top = tf2.margin_bottom = Emu(0)
    for i, line in enumerate(body):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(2); p.line_spacing = 1.06
        r = p.add_run(); r.text = line; r.font.size = Pt(12.5)
        r.font.color.rgb = SLATE; r.font.name = "Segoe UI"
    return panel


# ============================================================ 1. TITLE
s = slide(); bg(s, NAVY)
rect(s, Inches(0), Inches(0), SW, Inches(0.28), AZURE)
rect(s, Inches(8.7), Inches(0), Inches(4.63), SH, NAVY2)
for i, cx in enumerate([9.4, 10.15, 10.9, 11.65, 12.4]):
    for j, cy in enumerate([1.2, 1.95, 2.7, 3.45, 4.2, 4.95]):
        col = [AZURE, CYAN, TEAL][(i + j) % 3]
        rect(s, Inches(cx), Inches(cy), Inches(0.34), Inches(0.34), col, shape=MSO_SHAPE.OVAL)
txt(s, Inches(0.8), Inches(1.55), Inches(8), Inches(0.5),
    [[("AZURE COSMOS DB", 14, CYAN, True, False)]], space_after=0)
txt(s, Inches(0.75), Inches(2.0), Inches(8.2), Inches(2.0),
    [[("OmniVec", 72, WHITE, True, False)],
     [("Universal Vector Ingestion Platform", 26, RGBColor(0xB9,0xCD,0xEC), False, False)]],
    space_after=8)
rect(s, Inches(0.82), Inches(4.2), Inches(2.0), Inches(0.07), TEAL)
txt(s, Inches(0.8), Inches(4.45), Inches(7.7), Inches(1.4),
    [[("Turn any data into AI-ready vector search \u2014", 18, WHITE, False, False)],
     [("in minutes, on Azure. No pipeline to build.", 18, WHITE, False, False)]],
    space_after=4, line_spacing=1.15)
txt(s, Inches(0.8), Inches(6.45), Inches(9), Inches(0.5),
    [[("What it offers  \u00b7  Use cases  \u00b7  Why OmniVec", 13, RGBColor(0x8F,0xA8,0xCC), False, True)]],
    space_after=0)

# ============================================================ 2. PROBLEM
s = slide(); y = content_header(s, "Where projects stall", "Your customers want RAG. The pipeline is the hard part.", 2)
txt(s, Inches(0.72), y, Inches(11.9), Inches(0.7),
    [[("Everyone wants semantic search and grounded AI. Getting data reliably into a vector index is where weeks disappear \u2014 before any value ships.", 15, SLATE, False, False)]],
    line_spacing=1.15)
cy = Inches(2.6); cw = Inches(3.86); ch = Inches(1.95); gap = Inches(0.2)
card(s, Inches(0.72), cy, cw, ch, "Custom glue per source",
     ["A new connector for every", "Blob, Cosmos, Postgres, SQL.", "Brittle and unowned."], accent=AZURE)
card(s, Inches(0.72)+cw+gap, cy, cw, ch, "Model & GPU sprawl",
     ["Text, image, PDF & OCR each", "need different models, GPUs", "and API contracts."], accent=CYAN)
card(s, Inches(0.72)+2*(cw+gap), cy, cw, ch, "Data goes stale",
     ["Docs change constantly.", "Re-embedding & dedup logic", "gets reinvented each time."], accent=TEAL)
cy2 = cy + ch + Inches(0.2)
card(s, Inches(0.72), cy2, cw, ch, "Ops burden up front",
     ["Autoscaling, retries, health", "and back-pressure \u2014 before", "the first vector is searched."], accent=AMBER)
card(s, Inches(0.72)+cw+gap, cy2, cw, ch, "Slow to demo",
     ["Weeks of integration before", "you can show a customer", "anything working."], accent=AZURE)
card(s, Inches(0.72)+2*(cw+gap), cy2, cw, ch, "Security to re-solve",
     ["Identity, secrets, RBAC and", "network hardening land on", "the engineering team."], accent=CYAN)

# ============================================================ 3. WHAT IT OFFERS (HERO)
s = slide(); bg(s, NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), SH, AZURE)
txt(s, Inches(0.7), Inches(0.5), Inches(11), Inches(0.4),
    [[("THE PITCH IN ONE SLIDE", 13, CYAN, True, False)]], space_after=0)
txt(s, Inches(0.68), Inches(0.92), Inches(12.0), Inches(1.0),
    [[("Ship vector search \u2014 ", 34, WHITE, True, False),
      ("skip the pipeline.", 34, CYAN, True, False)]], space_after=0)
txt(s, Inches(0.72), Inches(1.7), Inches(11.8), Inches(0.7),
    [[("OmniVec gives your customers a production vector-ingestion platform out of the box, so their engineers focus on the app \u2014 not the plumbing.", 15.5, RGBColor(0xC5,0xD6,0xEF), False, False)]],
    line_spacing=1.2)
offers = [
    ("\u26A1", "Minutes, not months", ["One command stands up the whole", "platform. First vectors searchable", "the same day."], AZURE),
    ("</>", "Zero glue code", ["Configure sources, models &", "destinations. OmniVec handles", "extraction, chunking, retries, scale."], CYAN),
    ("\u25C8", "Works with your stack", ["Any source, any vector store,", "your model \u2014 Azure OpenAI or", "self-hosted GPU. No lock-in."], TEAL),
    ("\u2713", "Production-grade", ["Managed identity, RBAC, autoscaling,", "health & metrics \u2014 secure and", "observable from day one."], AMBER),
]
vw = Inches(2.92); vh = Inches(2.85); vg = Inches(0.14); x0 = Inches(0.72); vy = Inches(2.75)
for i, (g, t, b, ac) in enumerate(offers):
    x = x0 + i * (vw + vg)
    p = rect(s, x, vy, vw, vh, NAVY2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    p.adjustments[0] = 0.05
    chip = rect(s, x + Inches(0.24), vy + Inches(0.26), Inches(0.6), Inches(0.6), ac,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    chip.adjustments[0] = 0.25
    ctf = chip.text_frame; ctf.margin_top = ctf.margin_bottom = Emu(0)
    cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run(); cr.text = g; cr.font.size = Pt(20); cr.font.bold = True
    cr.font.color.rgb = WHITE; cr.font.name = "Segoe UI"
    txt(s, x + Inches(0.24), vy + Inches(1.0), vw - Inches(0.45), Inches(0.6),
        [[(t, 16, WHITE, True, False)]], space_after=0)
    bullets(s, x + Inches(0.24), vy + Inches(1.55), vw - Inches(0.42), Inches(1.2), b,
            size=11.5, color=RGBColor(0xC5,0xD6,0xEF), marker="", marker_color=ac, gap=1,
            line_spacing=1.12)
footer(s, 3)

# ============================================================ 4. USE CASES
s = slide(); y = content_header(s, "Use cases", "What your customers can build", 4)
uses = [
    ("\U0001F916", "RAG & copilots", ["Ground LLMs and chatbots", "in enterprise knowledge."], AZURE),
    ("\U0001F50D", "Semantic search", ["Search by meaning across", "docs, tickets and wikis."], CYAN),
    ("\U0001F4C4", "Multi-modal doc search", ["PDFs, scans and images", "via OCR and vision models."], TEAL),
    ("\U0001F504", "Always-fresh index", ["Change feed re-embeds", "as source data changes."], AMBER),
    ("\U0001F517", "Similarity & recs", ["Find related items,", "deduplicate and cluster."], AZURE),
    ("\U0001F5C3", "Vectorize what you have", ["Backfill databases & blobs", "into a searchable index."], CYAN),
]
cw = Inches(3.86); ch = Inches(2.0); gx = Inches(0.14); gy = Inches(0.2)
for i, (g, t, b, ac) in enumerate(uses):
    x = Inches(0.72) + (i % 3) * (cw + gx)
    yy = y + Inches(0.1) + (i // 3) * (ch + gy)
    icon_card(s, x, yy, cw, ch, g, t, b, ac)

# ============================================================ 5. WHAT'S IN THE BOX
s = slide(); y = content_header(s, "What's in the box", "Everything to go from data to vectors", 5)
feats = [
    ("Universal connectors", ["Blob, Cosmos DB, PostgreSQL,", "MSSQL in \u2192 vector stores out."], AZURE),
    ("Bring-your-own model", ["Azure OpenAI or self-hosted", "GPU (BGE, CLIP, DSE-Qwen2)."], CYAN),
    ("Smart chunking", ["Truncate or overlap-chunk,", "controlled per pipeline."], TEAL),
    ("Real-time sync", ["Cosmos change feed & Event", "Grid keep vectors current."], AMBER),
    ("Three interfaces", ["Web UI, CLI and REST API", "for every workflow."], AZURE),
    ("Elastic & cost-aware", ["Autoscaling workers;", "GPU scale-to-zero when idle."], CYAN),
]
cw = Inches(3.86); ch = Inches(1.9); gx = Inches(0.14); gy = Inches(0.2)
for i, (t, b, ac) in enumerate(feats):
    x = Inches(0.72) + (i % 3) * (cw + gx)
    yy = y + Inches(0.05) + (i // 3) * (ch + gy)
    card(s, x, yy, cw, ch, t, b, accent=ac, title_size=15, body_size=12.5)

# ============================================================ 6. WORKS WITH YOUR DATA
s = slide(); y = content_header(s, "Meets data where it lives", "Works with your data \u2014 no rip-and-replace", 6)
fy = y + Inches(0.05); fh = Inches(1.5)
stages = [("SOURCES", ["Azure Blob", "Cosmos DB", "PostgreSQL", "MSSQL"], AZURE),
          ("EMBED", ["Azure OpenAI", "BGE \u00b7 CLIP", "DSE-Qwen2"], CYAN),
          ("VECTOR STORES", ["Cosmos DB Vector", "pgvector", "MSSQL"], TEAL)]
bw = Inches(3.55); bgp = Inches(0.55); x0 = Inches(0.9)
for i, (t, items, ac) in enumerate(stages):
    x = x0 + i * (bw + bgp)
    p = rect(s, x, fy, bw, fh, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    p.adjustments[0] = 0.07
    rect(s, x, fy, bw, Inches(0.09), ac)
    txt(s, x, fy + Inches(0.2), bw, Inches(0.35), [[(t, 13, ac, True, False)]],
        align=PP_ALIGN.CENTER, space_after=0)
    txt(s, x, fy + Inches(0.55), bw, Inches(0.9),
        [[("  \u00b7  ".join(items), 13.5, INK, False, False)]], align=PP_ALIGN.CENTER,
        space_after=0, line_spacing=1.2)
    if i < 2:
        rect(s, x + bw + Inches(0.05), fy + Inches(0.58), Inches(0.45), Inches(0.34), ac,
             shape=MSO_SHAPE.CHEVRON)
ly = fy + fh + Inches(0.35)
bullets(s, Inches(0.9), ly, Inches(5.85), Inches(2.4), [
    ("Your model, your choice", " \u2014 external APIs or GPUs in your own cluster."),
    ("No data lock-in", " \u2014 vectors land in stores you already run."),
], size=15, gap=12)
bullets(s, Inches(7.1), ly, Inches(5.6), Inches(2.4), [
    ("Data can stay in-tenant", " \u2014 self-host models so content never leaves."),
    ("Mix and match", " \u2014 many sources & destinations per platform."),
], size=15, gap=12)

# ============================================================ 7. MULTI-MODAL + SMART
s = slide(); y = content_header(s, "More than text", "Multi-modal, and smart about content", 7)
half = Inches(5.9)
card(s, Inches(0.72), y, half, Inches(2.55), "Every modality", [], accent=AZURE)
bullets(s, Inches(1.05), y+Inches(0.7), half-Inches(0.6), Inches(1.7), [
    ("Text", " \u2014 short fields to long documents."),
    ("Images", " \u2014 CLIP embeddings for visual search."),
    ("PDF pages", " \u2014 vision models read layout directly."),
    ("Scanned PDFs", " \u2014 PaddleOCR extracts text, then embeds."),
], size=14, gap=9)
card(s, Inches(6.75), y, half, Inches(2.55), "Smart chunking", [], accent=CYAN)
bullets(s, Inches(7.08), y+Inches(0.7), half-Inches(0.6), Inches(1.7), [
    ("Truncate", " \u2014 one vector per document. Simplest."),
    ("Chunk", " \u2014 overlapping splits by chars or tokens."),
    ("Store text", " \u2014 keep chunk text alongside the vector."),
    ("Custom IDs", " \u2014 stable patterns for citations & updates."),
], size=14, gap=9)
by = y + Inches(2.8)
card(s, Inches(0.72), by, Inches(11.9), Inches(1.5), "Always up to date", [], accent=TEAL)
bullets(s, Inches(1.05), by+Inches(0.6), Inches(5.6), Inches(0.8), [
    ("Inline mode", " \u2014 embed in place, high throughput."),
], size=13.5, gap=8)
bullets(s, Inches(6.9), by+Inches(0.6), Inches(5.5), Inches(0.8), [
    ("Queue mode", " \u2014 durable, retryable, autoscaled."),
], size=13.5, gap=8)

# ============================================================ 8. DEPLOY IN ONE COMMAND
s = slide(); y = content_header(s, "Fastest path to value", "Deploy the whole platform in one command", 8)
cb = rect(s, Inches(0.72), y, Inches(11.9), Inches(1.1), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
cb.adjustments[0] = 0.12
txt(s, Inches(1.05), y+Inches(0.16), Inches(11.2), Inches(0.8),
    [[("$  azd up", 22, CYAN, True, False),
      ("      # provisions AKS, Cosmos DB, ACR, Key Vault, Storage, Service Bus, Event Grid \u2014 then deploys via Helm  (~15\u201325 min)", 12.5, RGBColor(0x9F,0xB6,0xD8), False, False)]],
    space_after=0, anchor=MSO_ANCHOR.MIDDLE)
ry = y + Inches(1.4)
val = [("Same-day demos", "Stand up a customer POC in ~20 min", AZURE),
       ("~$5\u201310/day default", "2\u00d7 B4ms, serverless Cosmos, no GPU", TEAL),
       ("Turn-key teardown", "azd down --purge stops all charges", CYAN),
       ("Repeatable", "Config saved as tags; redeploy anywhere", AMBER),
       ("UI, CLI & API", "Drive it however the team prefers", AZURE),
       ("Grows with them", "Add GPU models & sources later", TEAL)]
cw = Inches(3.86); ch = Inches(1.15); gx = Inches(0.14); gy = Inches(0.16)
for i, (t, d, ac) in enumerate(val):
    x = Inches(0.72) + (i % 3) * (cw + gx)
    yy = ry + (i // 3) * (ch + gy)
    card(s, x, yy, cw, ch, t, [d], accent=ac, title_size=14, body_size=11.5)

# ============================================================ 9. HOW IT WORKS (LIGHT)
s = slide(); y = content_header(s, "Under the hood (briefly)", "How it works", 9)
txt(s, Inches(0.72), y, Inches(11.9), Inches(0.55),
    [[("A control plane you configure, a smart embedding router, and elastic workers \u2014 all on AKS, backed by Cosmos DB.", 15, SLATE, False, False)]],
    line_spacing=1.15)
row = y + Inches(0.85); bh = Inches(1.5)
boxes = [("Your data", "Blob \u00b7 Cosmos \u00b7 SQL \u00b7 Postgres", AZURE),
         ("OmniVec", "Configure & orchestrate\nautoscaling workers", AZURE),
         ("DocGrok", "Routes to the right model\nAzure OpenAI or GPU", CYAN),
         ("Vector store", "Searchable embeddings\nin your database", TEAL)]
bw = Inches(2.72); gapb = Inches(0.28); x0 = Inches(0.9)
for i, (t, d, ac) in enumerate(boxes):
    x = x0 + i * (bw + gapb)
    p = rect(s, x, row, bw, bh, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    p.adjustments[0] = 0.08
    rect(s, x, row, bw, Inches(0.09), ac)
    txt(s, x + Inches(0.2), row + Inches(0.28), bw - Inches(0.35), Inches(0.5),
        [[(t, 16, INK, True, False)]], space_after=0)
    lines = d.split("\n")
    txt(s, x + Inches(0.2), row + Inches(0.72), bw - Inches(0.35), Inches(0.7),
        [[(ln, 12, SLATE, False, False)] for ln in lines], space_after=2, line_spacing=1.1)
    if i < 3:
        rect(s, x + bw + Inches(0.02), row + Inches(0.58), Inches(0.24), Inches(0.34), ac,
             shape=MSO_SHAPE.CHEVRON)
by = row + bh + Inches(0.4)
bullets(s, Inches(0.9), by, Inches(11.6), Inches(1.4), [
    ("Control plane", " \u2014 FastAPI API + a lightweight controller track sources, jobs & metrics in Cosmos DB."),
    ("Embedding router (DocGrok)", " \u2014 discovers models, runs PDF\u2192OCR\u2192text pipelines, parks idle GPUs."),
    ("Change capture", " \u2014 a .NET change-feed processor keeps embeddings in sync in real time."),
], size=14, gap=9)

# ============================================================ 10. PRODUCTION-GRADE
s = slide(); y = content_header(s, "Enterprise-ready", "Production-grade by default", 10)
sec = [
    ("Secure identity", ["Workload Identity federation \u2014", "no secrets stored in pods."], AZURE),
    ("Data stays yours", ["Self-host models so content", "never leaves your tenant."], TEAL),
    ("Fine-grained RBAC", ["Cosmos native RBAC; keys", "held in Azure Key Vault."], CYAN),
    ("Elastic scale", ["Workers 1\u201310; change feed", "matches Cosmos partitions."], AMBER),
    ("Observable", ["Health endpoints, per-pipeline", "metrics and UI dashboards."], AZURE),
    ("Threat-modelled", ["Documented STRIDE model,", "reviewed on a 6-month cycle."], TEAL),
]
cw = Inches(3.86); ch = Inches(1.75); gx = Inches(0.14); gy = Inches(0.18)
for i, (t, b, ac) in enumerate(sec):
    x = Inches(0.72) + (i % 3) * (cw + gx)
    yy = y + (i // 3) * (ch + gy)
    card(s, x, yy, cw, ch, t, b, accent=ac, title_size=14, body_size=12)

# ============================================================ 11. WHY OMNIVEC (BUILD VS BUY)
s = slide(); y = content_header(s, "Why OmniVec", "Build it yourself  vs.  OmniVec", 11)
half = Inches(5.85); gpc = Inches(0.2)
# left: build yourself
lx = Inches(0.72)
lp = rect(s, lx, y, half, Inches(4.35), LIGHT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
lp.adjustments[0] = 0.04
rect(s, lx, y, half, Inches(0.7), MUTE)
txt(s, lx, y + Inches(0.16), half, Inches(0.4),
    [[("BUILD IT YOURSELF", 15, WHITE, True, False)]], align=PP_ALIGN.CENTER, space_after=0)
bullets(s, lx + Inches(0.35), y + Inches(0.95), half - Inches(0.7), Inches(3.2), [
    "Weeks of custom glue code",
    "A new connector per data source",
    "Manual re-indexing when data changes",
    "DIY GPU hosting & autoscaling",
    "You own security, RBAC & networking",
    "Hard to demo anything early",
], size=14, color=INK, marker="\u2715", marker_color=ROSE, gap=13)
# right: omnivec
rx = lx + half + gpc
rp = rect(s, rx, y, half, Inches(4.35), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
rp.adjustments[0] = 0.04
rect(s, rx, y, half, Inches(0.7), AZURE)
txt(s, rx, y + Inches(0.16), half, Inches(0.4),
    [[("WITH OMNIVEC", 15, WHITE, True, False)]], align=PP_ALIGN.CENTER, space_after=0)
bullets(s, rx + Inches(0.35), y + Inches(0.95), half - Inches(0.7), Inches(3.2), [
    "Deploy the platform in ~20 minutes",
    "Universal connectors, no per-source code",
    "Real-time auto-sync via change feed",
    "Managed GPUs + scale-to-zero",
    "Secure & observable out of the box",
    "Customer demo the same day",
], size=14, color=WHITE, marker="\u2713", marker_color=TEAL, gap=13)

# ============================================================ 12. ROADMAP
s = slide(); bg(s, NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), SH, AZURE)
txt(s, Inches(0.7), Inches(0.5), Inches(11), Inches(0.4),
    [[("WHERE WE'RE HEADED", 13, CYAN, True, False)]], space_after=0)
txt(s, Inches(0.68), Inches(0.9), Inches(11.9), Inches(0.7),
    [[("Roadmap & future plans", 32, WHITE, True, False)]], space_after=0)
rect(s, Inches(0.72), Inches(1.6), Inches(1.4), Inches(0.06), CYAN)
road = [
    ("Hardening to GA", ["Move beyond public preview:", "stable APIs, fuller feature set."], AZURE),
    ("Network isolation", ["Private endpoints for Cosmos /", "Blob / Key Vault / AOAI + WAF."], CYAN),
    ("Zero-trust in-cluster", ["mTLS / service mesh and image", "signing (cosign) across AKS."], TEAL),
    ("Auto event provisioning", ["Blob Event Grid subscriptions", "wired up on pipeline create."], AMBER),
    ("More connectors", ["Broaden source & destination", "coverage across the data estate."], AZURE),
    ("Distribution & DX", ["Homebrew tap for the CLI;", "automated fuzz & E2E coverage."], CYAN),
]
cw = Inches(3.86); ch = Inches(1.9); gx = Inches(0.14); gy = Inches(0.2)
for i, (t, b, ac) in enumerate(road):
    x = Inches(0.72) + (i % 3) * (cw + gx)
    yy = Inches(2.05) + (i // 3) * (ch + gy)
    p = rect(s, x, yy, cw, ch, NAVY2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    p.adjustments[0] = 0.06
    rect(s, x, yy, cw, Inches(0.1), ac)
    txt(s, x+Inches(0.22), yy+Inches(0.24), cw-Inches(0.4), Inches(0.5),
        [[(t, 15, WHITE, True, False)]], space_after=0)
    bullets(s, x+Inches(0.22), yy+Inches(0.78), cw-Inches(0.38), Inches(1.0), b,
            size=11.5, color=RGBColor(0xC5,0xD6,0xEF), marker="", marker_color=ac, gap=2,
            line_spacing=1.12)
footer(s, 12)

# ============================================================ 13. CTA
s = slide(); bg(s, NAVY)
rect(s, Inches(0), Inches(0), SW, Inches(0.28), AZURE)
rect(s, Inches(0), Inches(7.22), SW, Inches(0.28), TEAL)
txt(s, Inches(0.8), Inches(1.0), Inches(11.7), Inches(0.5),
    [[("THE TAKEAWAY", 14, CYAN, True, False)]], space_after=0)
txt(s, Inches(0.78), Inches(1.5), Inches(11.8), Inches(1.6),
    [[("Any data source \u2192 embeddings \u2192 vector search,", 32, WHITE, True, False)],
     [("deployed on Azure in a single command.", 32, WHITE, True, False)]],
    space_after=4, line_spacing=1.08)
takeaways = [
    ("Universal", "One platform for every source, modality & destination."),
    ("Managed", "Change capture, chunking, retries & scaling handled for you."),
    ("Fast to prove", "Stand up a working customer demo in ~20 minutes."),
]
for i, (t, d) in enumerate(takeaways):
    yy = Inches(3.45) + i*Inches(0.72)
    rect(s, Inches(0.85), yy+Inches(0.06), Inches(0.24), Inches(0.24), [AZURE,CYAN,TEAL][i],
         shape=MSO_SHAPE.OVAL)
    txt(s, Inches(1.35), yy, Inches(11), Inches(0.6),
        [[(t + "  \u2014  ", 18, WHITE, True, False), (d, 16, RGBColor(0xC5,0xD6,0xEF), False, False)]],
        space_after=0)
rect(s, Inches(0.85), Inches(5.95), Inches(2.0), Inches(0.06), CYAN)
txt(s, Inches(0.82), Inches(6.15), Inches(11.6), Inches(0.6),
    [[("Get started:  ", 15, WHITE, True, False),
      ("github.com/AzureCosmosDB/OmniVec", 15, CYAN, True, False),
      ("   \u00b7   azd up", 15, RGBColor(0xC5,0xD6,0xEF), False, False)]],
    space_after=0)

prs.save("OmniVec-Overview.pptx")
print("Saved OmniVec-Overview.pptx with", len(prs.slides._sldIdLst), "slides")
